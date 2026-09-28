# -*- coding: utf-8 -*-
"""
VALIDADOR DE FOTOS POR CÓDIGO DE COR - BRASIL BOTÕES
"""

import os
import re
from datetime import datetime
from pathlib import Path

import openpyxl

# ======================= CONFIGURAÇÃO DE CAMINHOS DINÂMICOS =======================
BASE_DIR = Path(__file__).resolve().parent.parent

PASTA_FOTOS = BASE_DIR / "docs" / "imagens"
ARQUIVO_TABELA = BASE_DIR / "apoio" / "tabela_cores_final.xlsx"
RELATORIO_SAIDA = BASE_DIR / "apoio"
# ==============================================================

EXTENSOES_VALIDAS = {".jpg", ".jpeg", ".png", ".webp"}


def carregar_codigos(caminho_tabela):
    """Lê a coluna cor_codigo e cor_nome da tabela de cores."""
    if not os.path.exists(caminho_tabela):
        print(f"ERRO: Tabela não encontrada em {caminho_tabela}")
        return []

    wb = openpyxl.load_workbook(caminho_tabela, data_only=True)
    ws = wb["Cores"] if "Cores" in wb.sheetnames else wb.active

    headers = [str(c.value).strip() if c.value else "" for c in ws[1]]
    
    idx_codigo = headers.index("cor_codigo") if "cor_codigo" in headers else 0
    idx_nome = headers.index("cor_nome") if "cor_nome" in headers else 1
    idx_imagem = headers.index("imagem_origem") if "imagem_origem" in headers else None

    codigos = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        codigo = row[idx_codigo] if idx_codigo < len(row) else None
        if not codigo:
            continue
        
        nome = row[idx_nome] if idx_nome < len(row) else None
        ja_tinha = bool(row[idx_imagem]) if idx_imagem is not None and idx_imagem < len(row) else False

        codigos.append({
            "cor_codigo": str(codigo).strip(),
            "cor_nome": nome,
            "ja_tinha_imagem_origem": ja_tinha,
        })
    return codigos


def listar_arquivos_fotos(pasta):
    """Varre a pasta (e subpastas) e retorna a lista de nomes de arquivo."""
    arquivos = []
    if not pasta.is_dir():
        return arquivos
    for raiz, _subpastas, nomes in os.walk(pasta):
        for nome in nomes:
            ext = os.path.splitext(nome)[1].lower()
            if ext in EXTENSOES_VALIDAS:
                arquivos.append(nome)
    return arquivos


def encontrar_foto(codigo, arquivos):
    """Procura o código como um token isolado no nome do arquivo."""
    padrao = re.compile(
        r'(?:^|_)' + re.escape(codigo) + r'(?:_|\.)',
        re.IGNORECASE
    )
    for nome_arquivo in arquivos:
        if padrao.search(nome_arquivo):
            return nome_arquivo
    return None


def main():
    print(f"Lendo tabela de cores em: {ARQUIVO_TABELA}")
    codigos = carregar_codigos(ARQUIVO_TABELA)
    print(f"  -> {len(codigos)} códigos de cor na tabela")

    print(f"Varrendo pasta de fotos em: {PASTA_FOTOS}")
    arquivos = listar_arquivos_fotos(PASTA_FOTOS)
    print(f"  -> {len(arquivos)} imagens encontradas na pasta")

    print("Cruzando códigos com arquivos...")
    linhas_relatorio = []
    tem_foto_count = 0
    for item in codigos:
        arquivo_encontrado = encontrar_foto(item["cor_codigo"], arquivos)
        tem_foto = arquivo_encontrado is not None
        if tem_foto:
            tem_foto_count += 1
        linhas_relatorio.append({
            "cor_codigo": item["cor_codigo"],
            "cor_nome": item["cor_nome"],
            "tem_foto": "SIM" if tem_foto else "NAO",
            "arquivo_encontrado": arquivo_encontrado or "",
            "ja_tinha_imagem_origem_antes": "SIM" if item["ja_tinha_imagem_origem"] else "NAO",
        })

    wb_out = openpyxl.Workbook()
    ws_out = wb_out.active
    ws_out.title = "Relatorio"
    ws_out.append(["cor_codigo", "cor_nome", "tem_foto", "arquivo_encontrado", "ja_tinha_imagem_origem_antes"])
    for linha in linhas_relatorio:
        ws_out.append([
            linha["cor_codigo"],
            linha["cor_nome"],
            linha["tem_foto"],
            linha["arquivo_encontrado"],
            linha["ja_tinha_imagem_origem_antes"],
        ])
    for col in ["A", "B", "C", "D", "E"]:
        ws_out.column_dimensions[col].width = 30

    nome_relatorio = f"relatorio_fotos_cores_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx"
    caminho_relatorio = RELATORIO_SAIDA / nome_relatorio
    wb_out.save(caminho_relatorio)

    print()
    print("=" * 50)
    print(f"Total de códigos de cor: {len(codigos)}")
    print(f"Com foto encontrada: {tem_foto_count}")
    print(f"Sem foto encontrada: {len(codigos) - tem_foto_count}")
    print(f"Relatório salvo em: {caminho_relatorio}")
    print("=" * 50)


if __name__ == "__main__":
    main()