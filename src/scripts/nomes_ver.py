# -*- coding: utf-8 -*-
"""
EXTRATOR DE REPRESENTANTES E CLIENTES - BRASIL BOTOES
"""

import os
import json
from pathlib import Path
import openpyxl

# ======================= CONFIGURAÇÃO DE CAMINHOS =======================
BASE_DIR = Path(__file__).resolve().parent.parent.parent

PLANILHA_CLIENTES = BASE_DIR / "data" / "ESTOQUES PROJETO V2 - Clientes_teste.xlsm"
ABA_CLIENTES = "Cliente"
LINHA_CABECALHO = 11

PASTA_SAIDA = BASE_DIR / "apoio"
# ==============================================================


def limpar_texto(valor):
    if valor is None:
        return ""
    return " ".join(str(valor).strip().upper().split())


def processar_planilha():
    if not PLANILHA_CLIENTES.exists():
        print(f"ERRO: Planilha não encontrada no caminho:\n{PLANILHA_CLIENTES}")
        return

    print(f"Lendo planilha em: {PLANILHA_CLIENTES}")
    wb = openpyxl.load_workbook(PLANILHA_CLIENTES, read_only=True, data_only=True, keep_vba=False)

    if ABA_CLIENTES not in wb.sheetnames:
        print(f"ERRO: Aba '{ABA_CLIENTES}' não encontrada na planilha.")
        return

    ws = wb[ABA_CLIENTES]

    # Identificação dinâmica de colunas por cabeçalho (com fallback para colunas 12 e 13)
    headers = None
    for row in ws.iter_rows(min_row=LINHA_CABECALHO, max_row=LINHA_CABECALHO, values_only=True):
        headers = [str(h).strip().upper() if h else "" for h in row]
        break

    idx_rep = 12
    idx_cli = 13
    if headers:
        if "DESC" in headers:
            idx_rep = headers.index("DESC")
        elif "REPRESENTANTE" in headers:
            idx_rep = headers.index("REPRESENTANTE")

        if "OBS" in headers:
            idx_cli = headers.index("OBS")
        elif "CLIENTE" in headers:
            idx_cli = headers.index("CLIENTE")

    lista_linhas = []
    stats_rep = {}
    stats_cli = {}

    for num_linha, row in enumerate(
        ws.iter_rows(min_row=LINHA_CABECALHO + 1, values_only=True),
        start=LINHA_CABECALHO + 1
    ):
        rep_val = limpar_texto(row[idx_rep]) if len(row) > idx_rep else ""
        cli_val = limpar_texto(row[idx_cli]) if len(row) > idx_cli else ""

        if not rep_val and not cli_val:
            continue

        item = {"linha": num_linha}

        if rep_val:
            item["representante"] = rep_val
            stats_rep[rep_val] = stats_rep.get(rep_val, 0) + 1

        if cli_val:
            item["cliente"] = cli_val
            stats_cli[cli_val] = stats_cli.get(cli_val, 0) + 1

        lista_linhas.append(item)

    PASTA_SAIDA.mkdir(parents=True, exist_ok=True)

    caminho_dados = PASTA_SAIDA / "dados_linhas.json"
    with open(caminho_dados, "w", encoding="utf-8") as f:
        json.dump(lista_linhas, f, ensure_ascii=False, indent=2)

    relatorio_analise = {
        "resumo": {
            "total_linhas_com_dados": len(lista_linhas),
            "total_representantes_distintos": len(stats_rep),
            "total_clientes_distintos": len(stats_cli),
        },
        "representantes_encontrados": [
            {"nome": nome, "ocorrencias": count}
            for nome, count in sorted(stats_rep.items(), key=lambda x: x[1], reverse=True)
        ],
        "clientes_encontrados": [
            {"nome": nome, "ocorrencias": count}
            for nome, count in sorted(stats_cli.items(), key=lambda x: x[1], reverse=True)
        ],
    }

    caminho_analise = PASTA_SAIDA / "relatorio_analise.json"
    with open(caminho_analise, "w", encoding="utf-8") as f:
        json.dump(relatorio_analise, f, ensure_ascii=False, indent=2)

    print("\n" + "=" * 60)
    print("PROCESSAMENTO CONCLUÍDO!")
    print("=" * 60)
    print(f"Total de linhas extraídas: {len(lista_linhas)}")
    print(f"Representantes distintos: {len(stats_rep)}")
    print(f"Clientes distintos: {len(stats_cli)}")
    print("\nArquivos salvos:")
    print(f" -> {caminho_dados}")
    print(f" -> {caminho_analise}")


if __name__ == "__main__":
    processar_planilha()