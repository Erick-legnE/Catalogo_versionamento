"""
gerar_json_aprendiz.py — Catálogo Digital Brasil Botões

Lê:
  - data/ESTOQUES PROJETO V2 - Saldos2_teste.xlsm
  - apoio/tabela_cores_final.xlsx
  - docs/pastilhas.json (índice de pastilhas, quando existir)

Gera:
  - docs/produtos.json          <- arquivo usado pelo site
  - dist/produtos.json          <- espelho para compatibilidade com o projeto atual

Fotos:
  1) Procura primeiro uma FOTO REAL do produto em <raiz>/imagens/
  2) Se não houver, usa a PASTILHA da mesma cor como imagem ilustrativa
  3) Se não houver nenhum dos dois, deixa o produto sem imagem

Campos importantes adicionados ao JSON:
  - tem_foto:      True SOMENTE quando existe foto real do produto
  - tem_imagem:    True quando existe foto real OU pastilha
  - usa_pastilha:  True quando a imagem exibida é uma pastilha
  - tipo_imagem:   "produto", "pastilha" ou "sem_imagem"
  - aviso_imagem:  texto que a interface pode mostrar como badge/aviso

A ideia é não confundir uma pastilha de cor com uma foto real do botão.
"""

from __future__ import annotations

import json
import re
import shutil
from pathlib import Path
from typing import Iterable

import openpyxl


# ─── CAMINHOS ─────────────────────────────────────────────────────────────────

BASE_DIR = Path(__file__).resolve().parent.parent.parent

PLANILHA = BASE_DIR / "data" / "ESTOQUES PROJETO V2 - Saldos2_teste.xlsm"
TABELA_CORES = BASE_DIR / "apoio" / "tabela_cores_final.xlsx"

# Pasta onde ficam as fotos reais/originais dos produtos.
PASTA_FOTOS_ORIGEM = BASE_DIR / "imagens"

# A página catalogo_geral.html usa IMAGES_BASE = 'imagens/'.
# Por isso as imagens publicadas precisam estar dentro de docs/imagens/.
DOCS_DIR = BASE_DIR / "docs"
PASTA_IMAGENS_DOCS = DOCS_DIR / "imagens"
PASTA_PASTILHAS_DOCS = PASTA_IMAGENS_DOCS / "pastilhas"
ARQUIVO_PASTILHAS_JSON = DOCS_DIR / "pastilhas.json"
ARQUIVO_JSON_DOCS = DOCS_DIR / "produtos.json"

# Mantém o dist porque o projeto atual já usa essa pasta em alguns scripts.
DIST_DIR = BASE_DIR / "dist"
ARQUIVO_JSON_DIST = DIST_DIR / "produtos.json"

EXTENSOES_IMAGEM = (".jpg", ".jpeg", ".png", ".webp")


# ─── CORES ────────────────────────────────────────────────────────────────────

FAMILIAS_PALAVRAS_CHAVE = {
    "Branco":     ["branco"],
    "Bege":       ["bege", "creme", "marfim", "areia"],
    "Marrom":     ["marrom", "caramelo", "chocolate", "café", "cafe", "conhaque"],
    "Preto":      ["preto"],
    "Cinza":      ["cinza", "grafite", "chumbo"],
    "Azul":       ["azul", "marinho", "royal", "turquesa"],
    "Verde":      ["verde", "oliva", "militar"],
    "Amarelo":    ["amarelo", "canário", "canario", "mostarda"],
    "Vermelho":   ["vermelho", "vinho", "bordô", "bordo"],
    "Rosa":       ["rosa", "pink", "fúcsia", "fucsia"],
    "Laranja":    ["laranja"],
    "Roxo":       ["roxo", "lilás", "lilas", "violeta"],
    "Dourado":    ["dourado", "ouro"],
    "Prata":      ["prata", "prateado", "metálico", "metalico"],
    "Multicolor": ["multicolor", "mesclado", "mescla"],
    "Natural":    ["natural"],
}


def inferir_familia_por_nome(nome: object) -> str | None:
    nome_lower = str(nome or "").strip().lower()
    if not nome_lower:
        return None

    for familia, palavras in FAMILIAS_PALAVRAS_CHAVE.items():
        if any(palavra in nome_lower for palavra in palavras):
            return familia

    return None


def normalizar_cor(codigo: object) -> str:
    """Converte 3, '003', 'C003', '1005', 3.0 etc. para C003/C1005."""
    if codigo is None:
        return ""

    # Excel pode entregar números como int/float.
    if isinstance(codigo, (int, float)):
        try:
            numero = int(codigo)
            return f"C{numero:03d}" if numero < 1000 else f"C{numero}"
        except (TypeError, ValueError, OverflowError):
            pass

    texto = str(codigo).strip().upper()
    if not texto:
        return ""

    texto = re.sub(r"\s+", "", texto)
    if texto.startswith("C"):
        texto = texto[1:]

    # Trata strings vindas do Excel como "3.0".
    if re.fullmatch(r"\d+\.0+", texto):
        texto = texto.split(".", 1)[0]

    if texto.isdigit():
        numero = int(texto)
        return f"C{numero:03d}" if numero < 1000 else f"C{numero}"

    # Fallback conservador para códigos fora do padrão.
    return f"C{texto}"


# ─── HELPERS GERAIS ───────────────────────────────────────────────────────────


def montar_localizacao(corredor: object, caixa: object) -> str:
    corredor = str(corredor or "").strip()
    caixa = str(caixa or "").strip()

    if not corredor and not caixa:
        return ""

    partes: list[str] = []

    if corredor:
        partes.append(f"Corredor {corredor}")

    if caixa:
        try:
            caixa_fmt = f"{int(float(caixa)):03d}"
        except (TypeError, ValueError):
            caixa_fmt = caixa
        partes.append(f"Caixa {caixa_fmt}")

    return " / ".join(partes)


def montar_prefixo_imagem(modelo: object, cor_codigo: object, ting: object) -> str:
    modelo_txt = str(modelo or "").strip()
    cor = normalizar_cor(cor_codigo)
    ting_txt = str(ting or "").strip()

    if ting_txt:
        return f"{modelo_txt}_{cor}+{ting_txt}"
    return f"{modelo_txt}_{cor}"


def caminho_relativo_imagens(caminho: str) -> str:
    """
    Converte caminhos do pastilhas.json para o formato esperado por produtos.json.

    Exemplo:
      imagens/pastilhas/pastilha_C003.png
    vira:
      pastilhas/pastilha_C003.png

    catalogo_geral.html adiciona 'imagens/' na frente.
    """
    texto = str(caminho or "").strip().replace("\\", "/")
    texto = texto.lstrip("./")

    if texto.startswith("docs/"):
        texto = texto[len("docs/"):]

    if texto.startswith("imagens/"):
        texto = texto[len("imagens/"):]

    return texto


def ordenar_imagens(nomes: Iterable[str]) -> list[str]:
    def ordem(nome: str) -> tuple[int, str]:
        n = nome.lower()
        if "_frente" in n:
            prioridade = 1
        elif "_verso" in n:
            prioridade = 2
        elif "_lado" in n or "_lat" in n:
            prioridade = 3
        elif "_detalhe" in n:
            prioridade = 4
        else:
            # Arquivo sem sufixo costuma ser a frente/principal.
            prioridade = 0
        return prioridade, n

    return sorted(dict.fromkeys(nomes), key=ordem)


# ─── FOTOS REAIS ──────────────────────────────────────────────────────────────


def buscar_imagens_reais(prefixo_com_furacao: str) -> list[str]:
    """
    Procura fotos reais em <raiz>/imagens/.

    Se encontrar, copia para docs/imagens/ e retorna SOMENTE o nome do arquivo,
    pois o HTML já adiciona o prefixo 'imagens/'.
    """
    encontrados: list[str] = []

    if not PASTA_FOTOS_ORIGEM.exists():
        return encontrados

    prefixo_lower = prefixo_com_furacao.lower()

    for arquivo in PASTA_FOTOS_ORIGEM.iterdir():
        if not arquivo.is_file():
            continue

        if arquivo.suffix.lower() not in EXTENSOES_IMAGEM:
            continue

        nome_stem = arquivo.stem.lower()
        corresponde = (
            nome_stem == prefixo_lower
            or nome_stem.startswith(prefixo_lower + "_")
        )

        if not corresponde:
            continue

        PASTA_IMAGENS_DOCS.mkdir(parents=True, exist_ok=True)
        destino = PASTA_IMAGENS_DOCS / arquivo.name

        # copy2 também atualiza um arquivo antigo com a versão atual da origem.
        shutil.copy2(arquivo, destino)
        encontrados.append(arquivo.name)

    return ordenar_imagens(encontrados)


# ─── PASTILHAS ────────────────────────────────────────────────────────────────


def carregar_indice_pastilhas() -> dict[str, list[str]]:
    """
    Cria mapa:
      C003 -> ['pastilhas/pastilha_C003.png', ...]

    Prioridade:
      1) docs/pastilhas.json, que já é o índice usado pela página de pastilhas
      2) varredura de docs/imagens/pastilhas/ como fallback

    O índice JSON é aceito como fonte de verdade. Isso também facilita builds em
    cópias leves do projeto em que as imagens não foram incluídas no ZIP.
    """
    indice: dict[str, list[str]] = {}

    if ARQUIVO_PASTILHAS_JSON.exists():
        try:
            with ARQUIVO_PASTILHAS_JSON.open("r", encoding="utf-8") as f:
                dados = json.load(f)

            if isinstance(dados, list):
                for item in dados:
                    if not isinstance(item, dict):
                        continue

                    cor_codigo = normalizar_cor(item.get("cor_codigo"))
                    if not cor_codigo:
                        continue

                    imagens_raw = item.get("imagens") or []
                    if isinstance(imagens_raw, str):
                        imagens_raw = [imagens_raw]

                    imagens = [
                        caminho_relativo_imagens(img)
                        for img in imagens_raw
                        if str(img or "").strip()
                    ]

                    if imagens:
                        indice[cor_codigo] = ordenar_imagens(imagens)

        except (OSError, json.JSONDecodeError) as exc:
            print(f"  [AVISO] Não foi possível ler {ARQUIVO_PASTILHAS_JSON}: {exc}")

    # Complementa ou substitui caso não exista manifest para alguma cor.
    if PASTA_PASTILHAS_DOCS.exists():
        padrao = re.compile(r"^pastilha_(C\d+)(?:_(?:frente|verso|lado|lat|detalhe))?$", re.I)

        encontrados_por_cor: dict[str, list[str]] = {}

        for arquivo in PASTA_PASTILHAS_DOCS.iterdir():
            if not arquivo.is_file() or arquivo.suffix.lower() not in EXTENSOES_IMAGEM:
                continue

            match = padrao.match(arquivo.stem)
            if not match:
                continue

            cor_codigo = normalizar_cor(match.group(1))
            encontrados_por_cor.setdefault(cor_codigo, []).append(
                f"pastilhas/{arquivo.name}"
            )

        for cor_codigo, imagens in encontrados_por_cor.items():
            if cor_codigo not in indice:
                indice[cor_codigo] = ordenar_imagens(imagens)

    return indice


def buscar_pastilha_fallback(
    cor_raw: object,
    indice_pastilhas: dict[str, list[str]],
) -> list[str]:
    cor_codigo = normalizar_cor(cor_raw)
    if not cor_codigo:
        return []

    return list(indice_pastilhas.get(cor_codigo, []))


# ─── TABELA DE CORES ─────────────────────────────────────────────────────────


def carregar_tabela_cores() -> dict[str, dict[str, str]]:
    tabela: dict[str, dict[str, str]] = {}

    if not TABELA_CORES.exists():
        print(f"  [AVISO] tabela_cores_final.xlsx não encontrada em {TABELA_CORES}")
        return tabela

    wb = openpyxl.load_workbook(TABELA_CORES, read_only=True, data_only=True)
    ws = wb.active
    headers = None

    for row in ws.iter_rows(values_only=True):
        if headers is None:
            headers = [str(h).strip().lower() if h else "" for h in row]
            continue

        if not row or not row[0]:
            continue

        linha = dict(zip(headers, row))
        codigo = normalizar_cor(linha.get("cor_codigo", ""))

        if codigo:
            tabela[codigo] = {
                "nome": str(linha.get("cor_nome", "") or "").strip(),
                "familia": str(linha.get("cor_familia", "") or "").strip(),
                "hex": str(linha.get("cor_hex", "") or "").strip(),
            }

    wb.close()
    return tabela


# ─── SAÍDA ────────────────────────────────────────────────────────────────────


def gravar_json(produtos: list[dict]) -> None:
    for caminho in (ARQUIVO_JSON_DOCS, ARQUIVO_JSON_DIST):
        caminho.parent.mkdir(parents=True, exist_ok=True)
        with caminho.open("w", encoding="utf-8") as f:
            json.dump(produtos, f, ensure_ascii=False, indent=2)


def main() -> None:
    print("=" * 64)
    print("  Brasil Botões — Gerando catálogo geral")
    print("  Foto real > Pastilha da cor > Sem imagem")
    print("=" * 64)

    if not PLANILHA.exists():
        print(f"\n[ERRO] Planilha não encontrada:\n  {PLANILHA}")
        raise SystemExit(1)

    print("\n[1/5] Carregando planilha de estoques...")
    print(f"      {PLANILHA}")
    wb = openpyxl.load_workbook(
        PLANILHA,
        read_only=True,
        keep_vba=True,
        data_only=True,
    )

    print("\n[2/5] Carregando tabela de cores...")
    tabela_cores = carregar_tabela_cores()
    print(f"      {len(tabela_cores)} cores carregadas.")

    print("\n[3/5] Carregando índice de pastilhas...")
    indice_pastilhas = carregar_indice_pastilhas()
    print(f"      {len(indice_pastilhas)} cores com pastilha disponíveis.")

    PASTA_IMAGENS_DOCS.mkdir(parents=True, exist_ok=True)

    print("\n[4/5] Processando produtos e imagens...")

    ws_est = wb["Estoque"] if "Estoque" in wb.sheetnames else wb.active

    headers = None
    dados_encontrados = False
    produtos: list[dict] = []

    produtos_com_foto = 0
    produtos_com_pastilha = 0
    produtos_sem_imagem = 0
    imagens_reais_copiadas = 0
    inativos_ignorados = 0

    for row in ws_est.iter_rows(values_only=True):
        if not dados_encontrados:
            if row and row[0] == "SEQ.":
                headers = [str(h).strip() if h else "" for h in row]
                dados_encontrados = True
            continue

        if not row or not row[0] or not str(row[0]).strip().isdigit():
            continue

        linha = dict(zip(headers, row))

        cod_red = str(linha.get("CÓD. RED.", "") or "").strip()
        modelo = str(linha.get("MODELO", "") or "").strip()
        tamanho = str(linha.get("TAMANHO", "") or "").strip()
        cor_raw = linha.get("COR", "")
        ting = str(linha.get("TING.", "") or "").strip()
        desc_cor = str(linha.get("DESC. COR", "") or "").strip()
        furos = str(linha.get("FUROS", "") or "").strip()
        acabam = str(linha.get("ACABAM.", "") or "").strip()
        obs = str(linha.get("OBS", "") or "").strip()
        material = str(linha.get("MATERIAL", "") or "").strip()
        situacao = str(
            linha.get("SITUAÇÃO", "")
            or linha.get("SITUACAO", "")
            or "Ativo"
        ).strip()

        if situacao.lower() != "ativo":
            inativos_ignorados += 1
            continue

        saldo = linha.get("QTD SALDO")
        um = str(linha.get("UM", "") or "").strip()
        descr = str(linha.get("DESCR.", "") or "").strip()
        corredor = str(linha.get("CORR", "") or "").strip()
        caixa = linha.get("CX", "")

        try:
            qtd = float(saldo) if saldo is not None else 0.0
        except (TypeError, ValueError):
            qtd = 0.0

        cor_codigo = normalizar_cor(cor_raw)
        ref_tabela = tabela_cores.get(cor_codigo, {})

        if desc_cor:
            cor_nome = desc_cor
            cor_familia = (
                inferir_familia_por_nome(desc_cor)
                or ref_tabela.get("familia", "")
            )
        else:
            cor_nome = ref_tabela.get("nome", "")
            cor_familia = ref_tabela.get("familia", "")

        cor_hex = ref_tabela.get("hex", "")

        prefixo = montar_prefixo_imagem(modelo, cor_raw, ting)
        prefixo_furacao = f"{prefixo}_{furos}" if furos else prefixo

        # Sempre calcula a pastilha disponível. Mesmo quando existe foto real,
        # esse campo fica no JSON e pode servir como fallback no navegador.
        pastilha_fallback = buscar_pastilha_fallback(cor_raw, indice_pastilhas)

        fotos_reais = buscar_imagens_reais(prefixo_furacao)

        if fotos_reais:
            imagens = fotos_reais
            tipo_imagem = "produto"
            usa_pastilha = False
            tem_foto = True
            tem_imagem = True
            aviso_imagem = ""
            produtos_com_foto += 1
            imagens_reais_copiadas += len(fotos_reais)

        elif pastilha_fallback:
            imagens = pastilha_fallback
            tipo_imagem = "pastilha"
            usa_pastilha = True
            tem_foto = False
            tem_imagem = True
            aviso_imagem = "Amostra de cor"
            produtos_com_pastilha += 1

        else:
            imagens = []
            tipo_imagem = "sem_imagem"
            usa_pastilha = False
            tem_foto = False
            tem_imagem = False
            aviso_imagem = "Sem foto"
            produtos_sem_imagem += 1

        produto = {
            "reduzido": cod_red,
            "referencia": modelo,
            "descricao": descr,
            "tamanho": tamanho,
            "cor_codigo": cor_codigo,
            "cor_nome": cor_nome,
            "cor_familia": cor_familia,
            "cor_hex": cor_hex,
            "tingimento": ting,
            "furacao": furos,
            "acabamento": acabam,
            "material": material,
            "obs": obs,
            "saldo": qtd,
            "unidade": um,
            "localizacao": montar_localizacao(corredor, caixa),
            "imagem_base": prefixo_furacao,

            # Imagem que o catálogo deve exibir agora.
            "imagens": imagens,

            # Pastilha disponível para esta cor, mesmo se houver foto real.
            # Útil como fallback extra no front-end.
            "pastilha_fallback": pastilha_fallback,

            # IMPORTANTE: tem_foto agora significa FOTO REAL DO PRODUTO.
            "tem_foto": tem_foto,
            "tem_imagem": tem_imagem,
            "usa_pastilha": usa_pastilha,
            "tipo_imagem": tipo_imagem,
            "aviso_imagem": aviso_imagem,
        }

        produtos.append(produto)

    wb.close()

    print("\n[5/5] Gravando JSON...")
    gravar_json(produtos)

    print(f"      Site: {ARQUIVO_JSON_DOCS}")
    print(f"      Espelho: {ARQUIVO_JSON_DIST}")

    print("\n" + "=" * 64)
    print(f"  Concluído: {len(produtos)} produtos exportados")
    print(f"  - Com foto real do produto : {produtos_com_foto}")
    print(f"  - Usando pastilha da cor   : {produtos_com_pastilha}")
    print(f"  - Sem nenhuma imagem       : {produtos_sem_imagem}")
    print(f"  - Fotos reais copiadas     : {imagens_reais_copiadas}")
    print(f"  - Inativos ignorados       : {inativos_ignorados}")
    print("=" * 64 + "\n")


if __name__ == "__main__":
    main()
