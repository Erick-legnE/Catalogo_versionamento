# -*- coding: utf-8 -*-
"""
GERADOR DE CATALOGOS E SUB-SITES - BRASIL BOTOES
"""

import os
import re
import glob
import json
import secrets
import string
import unicodedata
from pathlib import Path

import openpyxl

# ======================= CONFIGURAÇÃO DE CAMINHOS DINÂMICOS =======================
BASE_DIR = Path(__file__).resolve().parent.parent.parent

PLANILHA_CLIENTES = BASE_DIR / "data" / "ESTOQUES PROJETO V2 - Clientes_teste.xlsm"
ABA_CLIENTES = "Cliente"
LINHA_CABECALHO = 11

PASTA_DIST = BASE_DIR / "docs"
PASTA_IMAGENS = BASE_DIR / "imagens"
TABELA_CORES = BASE_DIR / "apoio" / "tabela_cores_final.xlsx"
ARQUIVO_USUARIOS = PASTA_DIST / "usuarios.json"

# Busca templates em src/templates/ ou apoio/
DIRS_TEMPLATES = [
    BASE_DIR / "src" / "templates",
    BASE_DIR / "apoio"
]

LINK_CATALOGO_GERAL = "https://brasilbotoes.github.io/catalogo"
LIMITE_MINIMO_OCORRENCIAS = 1  # Ajustado para incluir testes com menos de 5 itens

PALAVRAS_IGNORAR_SLUG = {
    "LTDA", "EIRELI", "ME", "EPP", "SA", "S/A", "CIA", "COMERCIO",
    "COMERCIAL", "REPRESENTACOES", "REPRESENTACAO", "INDUSTRIA",
    "INDUSTRIAL", "DISTRIBUIDORA", "IMPORTACAO", "EXPORTACAO",
    "DE", "DO", "DA", "DOS", "DAS", "E", "S", "A"
}

EXTENSOES_VALIDAS = [".jpg", ".jpeg", ".png", ".webp"]
# ==============================================================


def encontrar_template(nomes_possiveis):
    for d in DIRS_TEMPLATES:
        for nome in nomes_possiveis:
            p = d / nome
            if p.exists():
                return p
    return None


def safe_get(row, idx):
    if idx is not None and 0 <= idx < len(row):
        return row[idx]
    return None


def limpar_texto(valor):
    if valor is None:
        return ""
    return " ".join(str(valor).strip().upper().split())


def gerar_slug_curto(texto):
    if not texto:
        return ""
    texto_nfkd = unicodedata.normalize('NFKD', str(texto))
    texto_ascii = texto_nfkd.encode('ascii', 'ignore').decode('utf-8').upper()
    palavras = re.findall(r'[A-Z0-9]+', texto_ascii)
    significativas = [p for p in palavras if p not in PALAVRAS_IGNORAR_SLUG]
    if not significativas:
        significativas = palavras
    return "-".join(significativas[:2]).lower()


def cor_formatada(cor_raw):
    try:
        return str(int(str(cor_raw).strip())).zfill(3)
    except (ValueError, TypeError):
        return str(cor_raw).strip() if cor_raw is not None else ""


def carregar_tabela_cores(caminho):
    cores = {}
    if not os.path.exists(caminho):
        return cores
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    ws = wb.active
    headers = None
    for row in ws.iter_rows(min_row=1, max_row=1, values_only=True):
        headers = [str(h).strip() if h else "" for h in row]
        break
    if not headers:
        return cores

    idx = {h: i for i, h in enumerate(headers)}
    for row in ws.iter_rows(min_row=2, values_only=True):
        codigo = safe_get(row, idx.get("cor_codigo"))
        if codigo is None:
            continue
        codigo_fmt = cor_formatada(codigo)
        cores[codigo_fmt] = {
            "nome": safe_get(row, idx.get("cor_nome")),
            "familia": safe_get(row, idx.get("cor_familia")),
            "hex": safe_get(row, idx.get("cor_hex")),
        }
    return cores


def buscar_foto(modelo, cor_fmt, ting, furos):
    if not PASTA_IMAGENS.is_dir():
        return None
    base = f"{modelo}_C{cor_fmt}"
    if ting:
        base += f"+{ting}"
    base += f"_{furos}"
    for ext in EXTENSOES_VALIDAS:
        candidato = PASTA_IMAGENS / (base + ext)
        if candidato.exists():
            return os.path.relpath(candidato, BASE_DIR).replace("\\", "/")
    padrao = str(PASTA_IMAGENS / (base + ".*"))
    achados = glob.glob(padrao)
    if achados:
        return os.path.relpath(achados[0], BASE_DIR).replace("\\", "/")
    return None


def ler_produtos():
    print(f"Lendo planilha de Clientes em: {PLANILHA_CLIENTES}")
    wb = openpyxl.load_workbook(PLANILHA_CLIENTES, read_only=True, data_only=True, keep_vba=False)
    ws = wb[ABA_CLIENTES]

    headers = None
    for row in ws.iter_rows(min_row=LINHA_CABECALHO, max_row=LINHA_CABECALHO, values_only=True):
        headers = [str(h).strip().upper() if h else "" for h in row]
        break

    def col(nomes):
        if not headers:
            return None
        for n in nomes:
            if n in headers:
                return headers.index(n)
        return None

    idx_codred = col(["CÓD. RED.", "COD. RED.", "CODRED"])
    idx_modelo = col(["MODELO", "REFERENCIA"])
    idx_tamanho = col(["TAMANHO"])
    idx_cor = col(["COR"])
    idx_ting = col(["TING.", "TING"])
    idx_desccor = col(["DESC. COR", "DESCCOR"])
    idx_furos = col(["FUROS"])
    idx_acabam = col(["ACABAM.", "ACABAMENTO"])
    idx_material = col(["MATERIAL"])
    idx_descr = col(["DESCR.", "DESCRICAO"])
    idx_qtd = col(["QTD SALDO", "SALDO"])
    idx_um = col(["UM"])
    
    # Busca coluna de cliente e representante com fallback para posições 12 e 13
    idx_rep = col(["REPRESENTANTE", "REP", "DESC"])
    if idx_rep is None:
        idx_rep = 12

    idx_cliente = col(["CLIENTE", "CLI", "OBS"])
    if idx_cliente is None:
        idx_cliente = 13

    cores = carregar_tabela_cores(TABELA_CORES)
    produtos_por_rep = {}
    produtos_por_cli = {}
    relatorio = []

    for row in ws.iter_rows(min_row=LINHA_CABECALHO + 1, values_only=True):
        codred = safe_get(row, idx_codred)
        rep_raw = safe_get(row, idx_rep) or ""
        cli_raw = safe_get(row, idx_cliente) or ""

        rep = limpar_texto(rep_raw)
        cli = limpar_texto(cli_raw)

        if codred is None or (not rep and not cli):
            continue

        modelo_val = safe_get(row, idx_modelo)
        modelo = str(modelo_val).strip() if modelo_val is not None else ""
        
        cor_raw = safe_get(row, idx_cor)
        cor_fmt = cor_formatada(cor_raw) if cor_raw is not None else ""
        
        ting_raw = safe_get(row, idx_ting)
        ting = str(ting_raw).strip() if ting_raw not in (None, "", "SIM") else None
        
        furos_val = safe_get(row, idx_furos)
        furos = str(furos_val).strip() if furos_val is not None else ""

        foto = buscar_foto(modelo, cor_fmt, ting, furos)
        info_cor = cores.get(cor_fmt, {})

        desc_cor_val = safe_get(row, idx_desccor)
        tamanho_val = safe_get(row, idx_tamanho)
        um_val = safe_get(row, idx_um)

        produto = {
            "cod_red": str(codred),
            "modelo": modelo,
            "tamanho": str(tamanho_val).strip() if tamanho_val is not None else None,
            "cor_codigo": f"C{cor_fmt}",
            "tingimento": ting,
            "cor_nome": desc_cor_val if desc_cor_val else info_cor.get("nome"),
            "cor_familia": info_cor.get("familia"),
            "cor_hex": info_cor.get("hex"),
            "furacao": furos,
            "acabamento": safe_get(row, idx_acabam),
            "material": safe_get(row, idx_material),
            "descricao": safe_get(row, idx_descr),
            "qtd_saldo": safe_get(row, idx_qtd),
            "unidade": um_val if um_val not in (None, "") else None,
            "cliente": cli,
            "representante": rep,
            "foto": foto,
        }

        if rep:
            produtos_por_rep.setdefault(rep, []).append(produto)
        if cli:
            produtos_por_cli.setdefault(cli, []).append(produto)

        relatorio.append({
            "representante": rep,
            "cliente": cli,
            "cod_red": produto["cod_red"],
            "modelo": modelo,
            "cor": produto["cor_codigo"],
            "foto_encontrada": "SIM" if foto else "NAO",
        })

    return produtos_por_rep, produtos_por_cli, relatorio


def gerar_senha_aleatoria(tamanho=8):
    caracteres = string.ascii_letters + string.digits
    return ''.join(secrets.choice(caracteres) for _ in range(tamanho))


def gerenciar_autenticacao(slugs_map, tipo_entidade="representante"):
    usuarios_dict = {}
    if ARQUIVO_USUARIOS.exists():
        try:
            with open(ARQUIVO_USUARIOS, "r", encoding="utf-8") as f:
                dados = json.load(f)
                for item in dados:
                    usuarios_dict[item["slug"]] = item
        except Exception:
            usuarios_dict = {}

    for nome, slug in sorted(slugs_map.items()):
        if slug in usuarios_dict:
            usuarios_dict[slug]["nome"] = nome
            usuarios_dict[slug]["tipo"] = tipo_entidade
        else:
            user_id = f"user_{len(usuarios_dict) + 1:03d}"
            usuarios_dict[slug] = {
                "id": user_id,
                "nome": nome,
                "slug": slug,
                "username": slug,
                "password": gerar_senha_aleatoria(8),
                "tipo": tipo_entidade,
                "ativo": True
            }

    lista_usuarios = list(usuarios_dict.values())
    ARQUIVO_USUARIOS.parent.mkdir(parents=True, exist_ok=True)
    with open(ARQUIVO_USUARIOS, "w", encoding="utf-8") as f:
        json.dump(lista_usuarios, f, ensure_ascii=False, indent=2)

    return usuarios_dict


def gerar_paginas(produtos_map, tipo_entidade="representantes"):
    if tipo_entidade == "clientes":
        path_template = encontrar_template(["template_cliente_index.html"])
    else:
        path_template = encontrar_template(["template_representante_index_2.html", "template_representante_index.html"])

    template_html = None
    if path_template and path_template.exists():
        print(f"Usando template: {path_template}")
        with open(path_template, "r", encoding="utf-8") as f:
            template_html = f.read()
    else:
        print(f"[AVISO CRÍTICO] Template para {tipo_entidade} não foi encontrado!")

    entidades_filtradas = {
        nome: prods for nome, prods in produtos_map.items() 
        if len(prods) >= LIMITE_MINIMO_OCORRENCIAS
    }

    slugs_map = {}
    slugs_usados = set()
    for nome in entidades_filtradas.keys():
        base = gerar_slug_curto(nome)
        slug = base
        c = 2
        while slug in slugs_usados:
            slug = f"{base}-{c}"
            c += 1
        slugs_usados.add(slug)
        slugs_map[nome] = slug

    mapa_usuarios = gerenciar_autenticacao(slugs_map, tipo_entidade)
    links = []

    for nome, produtos in entidades_filtradas.items():
        slug = slugs_map[nome]
        dados_auth = mapa_usuarios.get(slug, {})
        pasta_destino = PASTA_DIST / tipo_entidade / slug
        pasta_destino.mkdir(parents=True, exist_ok=True)

        produtos_subsite = []
        for p in produtos:
            p_copy = dict(p)
            if p_copy.get("foto"):
                p_copy["foto"] = f"../../{p_copy['foto']}"
            produtos_subsite.append(p_copy)

        payload_produtos = {
            "entidade_id": dados_auth.get("id"),
            "entidade_nome": nome,
            "tipo": tipo_entidade,
            "produtos": produtos_subsite
        }

        with open(pasta_destino / "produtos.json", "w", encoding="utf-8") as f:
            json.dump(payload_produtos, f, ensure_ascii=False, indent=2)

        if template_html:
            html = template_html.replace("{{REPRESENTANTE}}", nome)
            html = html.replace("{{CLIENTE}}", nome)
            html = html.replace("{{NOME_ENTIDADE}}", nome)
            html = html.replace("{{LINK_CATALOGO_GERAL}}", LINK_CATALOGO_GERAL)
            html = html.replace("{{REPRESENTANTE_ID}}", str(dados_auth.get("id", "")))
            
            with open(pasta_destino / "index.html", "w", encoding="utf-8") as f:
                f.write(html)
        else:
            print(f"  [ERRO] index.html não foi criado para '{nome}' por falta de template.")

        links.append((nome, slug, len(produtos)))
        print(f"  -> [{tipo_entidade.upper()}] {nome} ({slug}): {len(produtos)} produtos")

    return links


def exportar_indice_vitrine(links_rep, links_cli):
    vitrine = []
    
    for nome, slug, qtd in links_rep:
        vitrine.append({
            "nome": nome,
            "slug": slug,
            "tipo": "representante",
            "badge": "Representante",
            "qtd": qtd,
            "url": f"representantes/{slug}/index.html"
        })
        
    for nome, slug, qtd in links_cli:
        vitrine.append({
            "nome": nome,
            "slug": slug,
            "tipo": "cliente",
            "badge": "Cliente Especial",
            "qtd": qtd,
            "url": f"clientes/{slug}/index.html"
        })

    PASTA_DIST.mkdir(parents=True, exist_ok=True)
    caminho_vitrine = PASTA_DIST / "subsites_index.json"
    with open(caminho_vitrine, "w", encoding="utf-8") as f:
        json.dump(vitrine, f, ensure_ascii=False, indent=2)
    print(f"\nÍndice da vitrine gerado com sucesso em:\n -> {caminho_vitrine}")


def main():
    PASTA_DIST.mkdir(parents=True, exist_ok=True)
    prods_rep, prods_cli, relatorio = ler_produtos()
    
    print("\n--- GERANDO SUB-SITES DE REPRESENTANTES ---")
    links_rep = gerar_paginas(prods_rep, tipo_entidade="representantes")
    
    print("\n--- GERANDO SUB-SITES DE CLIENTES ---")
    links_cli = gerar_paginas(prods_cli, tipo_entidade="clientes")

    exportar_indice_vitrine(links_rep, links_cli)

    sem_foto = sum(1 for l in relatorio if l["foto_encontrada"] == "NAO")
    print(f"\nTotal de itens processados: {len(relatorio)}")
    print(f"Itens sem foto: {sem_foto}")


if __name__ == "__main__":
    main()